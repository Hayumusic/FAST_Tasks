from __future__ import annotations

import csv
import os
import sqlite3
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from flask import Flask, g, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("FAST_DB_PATH", BASE_DIR / "instance" / "fast_tracker.sqlite3")).expanduser()

NODE_TYPES = {"process", "subprocess", "task"}
USE_CASE_TYPES = {"Automation", "AI", "Agentic AI"}
STATUSES = {"Planned", "WIP", "Built", "Built and Deployed"}

app = Flask(__name__)
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


@app.after_request
def add_no_cache_headers(response):
    if request.path == "/" or request.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        DB_PATH.parent.mkdir(exist_ok=True)
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(_: BaseException | None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


def row_to_dict(row: sqlite3.Row) -> dict:
    return dict(row)


def init_db() -> None:
    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS nodes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            description TEXT NOT NULL DEFAULT '',
            type TEXT NOT NULL CHECK(type IN ('process', 'subprocess', 'task')),
            parent_id INTEGER REFERENCES nodes(id) ON DELETE CASCADE,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS pillars (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE,
            description TEXT NOT NULL DEFAULT ''
        );

        CREATE TABLE IF NOT EXISTS node_pillars (
            node_id INTEGER NOT NULL REFERENCES nodes(id) ON DELETE CASCADE,
            pillar_id INTEGER NOT NULL REFERENCES pillars(id) ON DELETE CASCADE,
            PRIMARY KEY (node_id, pillar_id)
        );

        CREATE TABLE IF NOT EXISTS use_cases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            node_id INTEGER NOT NULL REFERENCES nodes(id) ON DELETE CASCADE,
            title TEXT NOT NULL,
            problem_statement TEXT NOT NULL,
            use_case_type TEXT NOT NULL,
            short_description TEXT NOT NULL,
            status TEXT NOT NULL,
            contributor TEXT NOT NULL,
            created_date TEXT NOT NULL,
            updated_date TEXT NOT NULL,
            effort_hours REAL NOT NULL DEFAULT 0,
            cost_per_run REAL NOT NULL DEFAULT 0,
            frequency_per_month REAL NOT NULL DEFAULT 0,
            savings_percent REAL NOT NULL DEFAULT 0
        );
        """
    )
    ensure_columns(
        db,
        "use_cases",
        {
            "effort_hours": "REAL NOT NULL DEFAULT 0",
            "cost_per_run": "REAL NOT NULL DEFAULT 0",
            "frequency_per_month": "REAL NOT NULL DEFAULT 0",
            "savings_percent": "REAL NOT NULL DEFAULT 0",
        },
    )
    db.commit()


def ensure_columns(db: sqlite3.Connection, table: str, columns: dict[str, str]) -> None:
    existing = {row["name"] for row in db.execute(f"PRAGMA table_info({table})")}
    for name, definition in columns.items():
        if name not in existing:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {definition}")


def ids_arg(name: str) -> list[int]:
    raw = request.form.get(name) or request.json.get(name, "") if request.is_json else request.form.get(name, "")
    if isinstance(raw, list):
        return [int(x) for x in raw if str(x).strip()]
    return [int(x) for x in str(raw).split(",") if x.strip()]


def node_path_cte() -> str:
    return """
        WITH node_path AS (
            SELECT
                n.id AS node_id,
                CASE
                    WHEN n.type = 'process' THEN n.id
                    WHEN n.type = 'subprocess' THEN parent.id
                    WHEN n.type = 'task' THEN grandparent.id
                END AS process_id,
                CASE
                    WHEN n.type = 'process' THEN n.name
                    WHEN n.type = 'subprocess' THEN COALESCE(parent.name, '')
                    WHEN n.type = 'task' THEN COALESCE(grandparent.name, '')
                END AS process,
                CASE
                    WHEN n.type = 'subprocess' THEN n.id
                    WHEN n.type = 'task' THEN parent.id
                END AS subprocess_id,
                CASE
                    WHEN n.type = 'subprocess' THEN n.name
                    WHEN n.type = 'task' THEN COALESCE(parent.name, '')
                    ELSE ''
                END AS subprocess,
                CASE WHEN n.type = 'task' THEN n.id END AS task_id,
                CASE WHEN n.type = 'task' THEN n.name ELSE '' END AS task,
                n.type AS attached_level
            FROM nodes n
            LEFT JOIN nodes parent ON n.parent_id = parent.id
            LEFT JOIN nodes grandparent ON parent.parent_id = grandparent.id
        )
    """


def filtered_use_cases() -> list[dict]:
    params = request.args
    where = []
    values: list[str] = []
    if params.get("pillar"):
        where.append(
            "EXISTS (SELECT 1 FROM node_pillars np JOIN pillars p ON p.id = np.pillar_id WHERE np.node_id = uc.node_id AND p.name = ?)"
        )
        values.append(params["pillar"])
    for key, column in [("process", "path.process_id"), ("subprocess", "path.subprocess_id"), ("type", "uc.use_case_type"), ("status", "uc.status"), ("contributor", "uc.contributor")]:
        if params.get(key):
            where.append(f"{column} = ?")
            values.append(params[key])
    clause = "WHERE " + " AND ".join(where) if where else ""
    rows = get_db().execute(
        node_path_cte()
        + f"""
        SELECT
            uc.id,
            uc.node_id,
            GROUP_CONCAT(p.name, ', ') AS pillar,
            path.process_id,
            path.process,
            path.subprocess_id,
            path.subprocess,
            path.task_id,
            path.task,
            path.attached_level,
            uc.title,
            uc.problem_statement,
            uc.use_case_type,
            uc.short_description,
            uc.status,
            uc.contributor,
            uc.created_date,
            uc.updated_date,
            uc.effort_hours,
            uc.cost_per_run,
            uc.frequency_per_month,
            uc.savings_percent,
            ROUND(uc.cost_per_run * uc.frequency_per_month, 2) AS estimated_monthly_cost,
            ROUND(uc.cost_per_run * uc.frequency_per_month * uc.savings_percent / 100.0, 2) AS estimated_monthly_savings
        FROM use_cases uc
        JOIN node_path path ON path.node_id = uc.node_id
        LEFT JOIN node_pillars np ON np.node_id = uc.node_id
        LEFT JOIN pillars p ON p.id = np.pillar_id
        {clause}
        GROUP BY uc.id
        ORDER BY uc.updated_date DESC, uc.id DESC
        """,
        values,
    ).fetchall()
    return [row_to_dict(row) for row in rows]


def get_nodes() -> list[dict]:
    db = get_db()
    rows = db.execute(
        """
        SELECT n.*, COALESCE(GROUP_CONCAT(p.name, ', '), '') AS pillars,
               COALESCE(GROUP_CONCAT(p.id), '') AS pillar_ids
        FROM nodes n
        LEFT JOIN node_pillars np ON np.node_id = n.id
        LEFT JOIN pillars p ON p.id = np.pillar_id
        GROUP BY n.id
        ORDER BY CASE n.type WHEN 'process' THEN 1 WHEN 'subprocess' THEN 2 ELSE 3 END, n.name
        """
    ).fetchall()
    return [row_to_dict(row) for row in rows]


def find_or_create_node(db: sqlite3.Connection, name: str, node_type: str, parent_id: int | None) -> int:
    row = db.execute(
        "SELECT id FROM nodes WHERE lower(name) = lower(?) AND type = ? AND COALESCE(parent_id, 0) = COALESCE(?, 0)",
        (name, node_type, parent_id),
    ).fetchone()
    if row:
        return row["id"]
    cur = db.execute(
        "INSERT INTO nodes (name, type, parent_id, updated_at) VALUES (?, ?, ?, CURRENT_TIMESTAMP)",
        (name, node_type, parent_id),
    )
    return cur.lastrowid


def normalize_header(value: object) -> str:
    return "".join(ch for ch in str(value or "").strip().lower() if ch.isalnum())


def worksheet_rows(worksheet) -> list[dict]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell or "").strip() for cell in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(cell not in (None, "") for cell in row)]


def row_has(row: dict, aliases: set[str]) -> bool:
    headers = {normalize_header(key) for key in row}
    return bool(headers & aliases)


def parse_upload(file_storage) -> tuple[list[dict], list[dict]]:
    filename = (file_storage.filename or "").lower()
    raw = file_storage.read()
    if filename.endswith(".csv"):
        text = raw.decode("utf-8-sig")
        return list(csv.DictReader(text.splitlines(), restkey="_extra")), []
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    hierarchy_rows: list[dict] = []
    pillar_rows: list[dict] = []
    for worksheet in workbook.worksheets:
        rows = worksheet_rows(worksheet)
        if not rows:
            continue
        sample = rows[0]
        has_process = row_has(sample, {"process", "processname", "level1"})
        has_subprocess = row_has(sample, {"subprocess", "subprocessname", "subprocesses", "level2"})
        has_pillar = row_has(sample, {"pillar", "pillars", "pillarname", "name"})
        if has_process and has_subprocess:
            hierarchy_rows.extend(rows)
        elif has_pillar:
            pillar_rows.extend(rows)
    return hierarchy_rows, pillar_rows


def value_for(row: dict, aliases: set[str]) -> str:
    normalized = {normalize_header(key): value for key, value in row.items()}
    for alias in aliases:
        value = normalized.get(alias)
        if value not in (None, ""):
            return str(value).strip()
    return ""


def task_name_for(row: dict) -> str:
    return value_for(row, {"task", "tasks", "taskname", "level3"})


def pillar_names_for(row: dict) -> list[str]:
    raw = value_for(row, {"pillar", "pillars", "pillarname"})
    return [name.strip() for name in raw.replace("\n", ",").split(",") if name.strip()]


def save_pillar(db: sqlite3.Connection, name: str, description: str = "") -> int:
    clean_name = name.strip()
    clean_description = description.strip()
    db.execute("INSERT OR IGNORE INTO pillars (name, description) VALUES (?, ?)", (clean_name, clean_description))
    if clean_description:
        db.execute("UPDATE pillars SET description = ? WHERE name = ?", (clean_description, clean_name))
    return db.execute("SELECT id FROM pillars WHERE name = ?", (clean_name,)).fetchone()["id"]


def import_pillars(db: sqlite3.Connection, rows: list[dict]) -> dict[str, int]:
    imported: dict[str, int] = {}
    for row in rows:
        name = value_for(row, {"pillar", "pillars", "pillarname", "name"})
        if not name:
            continue
        description = value_for(row, {"description", "desc", "details"})
        imported[name] = save_pillar(db, name, description)
    return imported


def hierarchy_payload() -> list[dict]:
    nodes = get_nodes()
    use_cases = filtered_use_cases()
    counts = {
        node["id"]: {"total": 0, "Planned": 0, "WIP": 0, "Built": 0, "Built and Deployed": 0}
        for node in nodes
    }
    for case in use_cases:
        current = case["node_id"]
        while current:
            bucket = counts[current]
            bucket["total"] += 1
            bucket[case["status"]] += 1
            current = next((n["parent_id"] for n in nodes if n["id"] == current), None)
    by_parent: dict[int | None, list[dict]] = {}
    for node in nodes:
        node["counts"] = counts[node["id"]]
        node["children"] = []
        by_parent.setdefault(node["parent_id"], []).append(node)
    for node in nodes:
        node["children"] = by_parent.get(node["id"], [])
    return by_parent.get(None, [])


@app.route("/")
def index():
    init_db()
    asset_version = max((BASE_DIR / "static" / "app.css").stat().st_mtime_ns, (BASE_DIR / "static" / "app.js").stat().st_mtime_ns)
    return render_template("index.html", statuses=sorted(STATUSES), use_case_types=sorted(USE_CASE_TYPES), asset_version=asset_version)


@app.route("/healthz")
def healthz():
    init_db()
    get_db().execute("SELECT 1").fetchone()
    return jsonify({"ok": True})


@app.route("/api/bootstrap")
def bootstrap():
    init_db()
    use_cases = filtered_use_cases()
    return jsonify(
        {
            "nodes": get_nodes(),
            "pillars": [row_to_dict(r) for r in get_db().execute("SELECT * FROM pillars ORDER BY name")],
            "useCases": use_cases,
            "hierarchy": hierarchy_payload(),
        }
    )


@app.route("/api/nodes", methods=["POST"])
def create_node():
    data = request.get_json()
    node_type = data.get("type")
    if node_type not in NODE_TYPES:
        return jsonify({"error": "Invalid node type"}), 400
    db = get_db()
    cur = db.execute(
        "INSERT INTO nodes (name, description, type, parent_id, updated_at) VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)",
        (data.get("name", "").strip(), data.get("description", "").strip(), node_type, data.get("parent_id") or None),
    )
    node_id = cur.lastrowid
    save_node_pillars(db, node_id, data.get("pillar_ids", []))
    db.commit()
    return jsonify({"id": node_id})


@app.route("/api/hierarchy/import", methods=["POST"])
def import_hierarchy():
    init_db()
    uploaded = request.files.get("file")
    if not uploaded:
        return jsonify({"error": "Upload a CSV or Excel file"}), 400
    hierarchy_rows, pillar_rows = parse_upload(uploaded)
    db = get_db()
    if request.form.get("replace") == "true":
        db.execute("DELETE FROM use_cases")
        db.execute("DELETE FROM nodes")
        db.execute("DELETE FROM pillars")
    imported_pillars = import_pillars(db, pillar_rows)
    touched_processes: set[int] = set()
    touched_subprocesses: set[int] = set()
    task_count = 0
    for row in hierarchy_rows:
        process_name = value_for(row, {"process", "processname", "level1"})
        subprocess_name = value_for(row, {"subprocess", "subprocessname", "subprocesses", "level2"})
        task_name = task_name_for(row)
        if not process_name or not subprocess_name:
            continue
        process_id = find_or_create_node(db, process_name, "process", None)
        subprocess_id = find_or_create_node(db, subprocess_name, "subprocess", process_id)
        pillar_ids = [
            imported_pillars.get(name) or save_pillar(db, name)
            for name in pillar_names_for(row)
        ]
        if pillar_ids:
            save_node_pillars(db, process_id, pillar_ids)
            save_node_pillars(db, subprocess_id, pillar_ids)
        touched_processes.add(process_id)
        touched_subprocesses.add(subprocess_id)
        if task_name:
            task_id = find_or_create_node(db, task_name, "task", subprocess_id)
            if pillar_ids:
                save_node_pillars(db, task_id, pillar_ids)
            task_count += 1
    db.commit()
    return jsonify(
        {
            "ok": True,
            "processes": len(touched_processes),
            "subprocesses": len(touched_subprocesses),
            "tasks": task_count,
            "pillars": len(imported_pillars),
        }
    )


@app.route("/api/nodes/<int:node_id>", methods=["PUT", "DELETE"])
def node_detail(node_id: int):
    db = get_db()
    if request.method == "DELETE":
        db.execute("DELETE FROM nodes WHERE id = ?", (node_id,))
        db.commit()
        return jsonify({"ok": True})
    data = request.get_json()
    db.execute(
        "UPDATE nodes SET name = ?, description = ?, parent_id = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
        (data.get("name", "").strip(), data.get("description", "").strip(), data.get("parent_id") or None, node_id),
    )
    save_node_pillars(db, node_id, data.get("pillar_ids", []))
    db.commit()
    return jsonify({"ok": True})


def save_node_pillars(db: sqlite3.Connection, node_id: int, pillar_ids: list[int]) -> None:
    db.execute("DELETE FROM node_pillars WHERE node_id = ?", (node_id,))
    valid_ids = [int(pid) for pid in pillar_ids if str(pid).strip()]
    db.executemany("INSERT OR IGNORE INTO node_pillars (node_id, pillar_id) VALUES (?, ?)", [(node_id, pid) for pid in valid_ids])


@app.route("/api/pillars", methods=["POST"])
def create_pillar():
    data = request.get_json()
    cur = get_db().execute("INSERT INTO pillars (name, description) VALUES (?, ?)", (data.get("name", "").strip(), data.get("description", "").strip()))
    get_db().commit()
    return jsonify({"id": cur.lastrowid})


@app.route("/api/pillars/<int:pillar_id>", methods=["PUT", "DELETE"])
def pillar_detail(pillar_id: int):
    db = get_db()
    if request.method == "DELETE":
        db.execute("DELETE FROM pillars WHERE id = ?", (pillar_id,))
    else:
        data = request.get_json()
        db.execute("UPDATE pillars SET name = ?, description = ? WHERE id = ?", (data.get("name", "").strip(), data.get("description", "").strip(), pillar_id))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/use-cases", methods=["POST"])
def create_use_case():
    data = request.get_json()
    today = date.today().isoformat()
    db = get_db()
    cur = db.execute(
        """
        INSERT INTO use_cases
        (node_id, title, problem_statement, use_case_type, short_description, status, contributor, created_date, updated_date,
         effort_hours, cost_per_run, frequency_per_month, savings_percent)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            data.get("node_id"),
            data.get("title", "").strip(),
            data.get("problem_statement", "").strip(),
            data.get("use_case_type"),
            data.get("short_description", "").strip(),
            data.get("status"),
            data.get("contributor", "").strip(),
            data.get("created_date") or today,
            data.get("updated_date") or today,
            float(data.get("effort_hours") or 0),
            float(data.get("cost_per_run") or 0),
            float(data.get("frequency_per_month") or 0),
            float(data.get("savings_percent") or 0),
        ),
    )
    db.commit()
    return jsonify({"id": cur.lastrowid})


@app.route("/api/use-cases/<int:case_id>", methods=["PUT", "DELETE"])
def use_case_detail(case_id: int):
    db = get_db()
    if request.method == "DELETE":
        db.execute("DELETE FROM use_cases WHERE id = ?", (case_id,))
        db.commit()
        return jsonify({"ok": True})
    data = request.get_json()
    db.execute(
        """
        UPDATE use_cases
        SET node_id = ?, title = ?, problem_statement = ?, use_case_type = ?, short_description = ?,
            status = ?, contributor = ?, created_date = ?, updated_date = ?,
            effort_hours = ?, cost_per_run = ?, frequency_per_month = ?, savings_percent = ?
        WHERE id = ?
        """,
        (
            data.get("node_id"),
            data.get("title", "").strip(),
            data.get("problem_statement", "").strip(),
            data.get("use_case_type"),
            data.get("short_description", "").strip(),
            data.get("status"),
            data.get("contributor", "").strip(),
            data.get("created_date"),
            data.get("updated_date") or date.today().isoformat(),
            float(data.get("effort_hours") or 0),
            float(data.get("cost_per_run") or 0),
            float(data.get("frequency_per_month") or 0),
            float(data.get("savings_percent") or 0),
            case_id,
        ),
    )
    db.commit()
    return jsonify({"ok": True})


@app.route("/export")
def export_excel():
    init_db()
    rows = filtered_use_cases()
    wb = Workbook()
    ws = wb.active
    ws.title = "Use Cases"
    headers = [
        "Pillar",
        "Process",
        "Sub-process",
        "Task",
        "Use Case Title",
        "Current Problem / Opportunity",
        "Type",
        "Status",
        "Short Description",
        "Contributor",
        "Created Date",
        "Last Updated Date",
        "Attached Level",
        "Current Effort Spent / Run",
        "Equivalent Current Cost / Run",
        "Frequency / Month",
        "Savings From Automation %",
        "Current Monthly Cost",
        "Monthly Savings From Automation",
    ]
    ws.append(headers)
    for row in rows:
        ws.append(
            [
                row["pillar"],
                row["process"],
                row["subprocess"],
                row["task"],
                row["title"],
                row["problem_statement"],
                row["use_case_type"],
                row["status"],
                row["short_description"],
                row["contributor"],
                row["created_date"],
                row["updated_date"],
                row["attached_level"],
                row["effort_hours"],
                row["cost_per_run"],
                row["frequency_per_month"],
                row["savings_percent"],
                row["estimated_monthly_cost"],
                row["estimated_monthly_savings"],
            ]
        )
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20323A")
    ws.auto_filter.ref = ws.dimensions
    table = Table(displayName="AI_Automation_Use_Cases", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(44, max(14, max(len(str(ws.cell(row=r, column=col).value or "")) for r in range(1, ws.max_row + 1)) + 2))
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"AI_Automation_Use_Cases_{date.today().isoformat()}.xlsx"
    return send_file(buffer, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    with app.app_context():
        init_db()
    port = int(os.environ.get("PORT", "5001"))
    debug = os.environ.get("FLASK_DEBUG", "true").lower() == "true"
    app.run(debug=debug, host="0.0.0.0", port=port)
